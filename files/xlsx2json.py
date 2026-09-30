# -*- coding: utf-8 -*-
"""
Ядро конвертации xlsx -> json.
Консольный запуск (один раз, без слежения):
    python xlsx2json.py file.xlsx
Результат: <имя_xlsx>.json рядом с исходником.
"""
import sys
sys.dont_write_bytecode = True

import os
import re
import json
from datetime import datetime, date, time
from openpyxl import load_workbook


REF_FORMULA = re.compile(r'^\s*=\s*\$?[A-Z]{1,3}\$?\d+\s*$')

CONVERTER_NAME = "xlsx2json"
CONVERTER_VERSION = "1.0"


def is_ref_formula(f):
    return isinstance(f, str) and bool(REF_FORMULA.match(f))


def is_error(v):
    return isinstance(v, str) and len(v) >= 2 and v.startswith("#") and v.endswith("!")


def col_letter(idx):
    s = ""
    while idx > 0:
        idx, r = divmod(idx - 1, 26)
        s = chr(65 + r) + s
    return s


def format_value(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, time):
        return v.strftime("%H:%M")
    return v


def type_name(v):
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "float"
    if isinstance(v, (datetime, date)):
        return "date"
    if isinstance(v, time):
        return "time"
    if isinstance(v, str):
        return "error" if is_error(v) else "str"
    return type(v).__name__


def md_escape(s):
    if s is None:
        return ""
    s = str(s)
    s = s.replace("\\", "\\\\").replace("|", "\\|")
    s = s.replace("\n", " ").replace("\r", " ")
    return s


def short_sample(v, limit=60):
    s = str(v)
    if len(s) > limit:
        s = s[: limit - 3] + "..."
    return s


def convert(xlsx_path, want_report=True):
    """
    Возвращает (result, errors, report).
    result  — {sheet_name: [ {col: val}, ... ]}
    errors  — список строк-замечаний
    report  — markdown-текст (или "" если want_report=False)
    """
    wb_vals = load_workbook(xlsx_path, data_only=True)
    wb_frms = load_workbook(xlsx_path, data_only=False)

    result = {}
    errors = []
    report_lines = []

    for sheet_name in wb_vals.sheetnames:
        ws_v = wb_vals[sheet_name]
        ws_f = wb_frms[sheet_name]

        all_rows = list(ws_v.iter_rows())
        if not all_rows:
            result[sheet_name] = []
            continue

        header_cells = all_rows[0]
        n_cols = len(header_cells)

        headers = []
        for cell in header_cells:
            h = cell.value
            if h is None or (isinstance(h, str) and h.strip() == ""):
                headers.append(None)
            else:
                headers.append(str(h).strip())

        for ci in range(n_cols):
            if headers[ci] is not None:
                continue
            has_data = False
            for r in all_rows[1:]:
                if ci >= len(r):
                    continue
                v = r[ci].value
                if v is None:
                    continue
                if isinstance(v, str) and v.strip() == "":
                    continue
                f = ws_f.cell(row=r[ci].row, column=ci + 1).value
                if is_ref_formula(f):
                    continue
                has_data = True
                break
            if has_data:
                letter = col_letter(ci + 1)
                headers[ci] = f"Колонка {letter}"
                errors.append(
                    f"Лист '{sheet_name}': пустой заголовок в столбце {letter}, "
                    f"использовано имя 'Колонка {letter}'"
                )

        col_types = {}
        col_samples = {}
        data_rows = []

        for r in all_rows[1:]:
            row_dict = {}
            row_empty = True
            for ci in range(n_cols):
                if ci >= len(r):
                    break
                h = headers[ci]
                if h is None:
                    continue
                cell = r[ci]
                f = ws_f.cell(row=cell.row, column=ci + 1).value
                if is_ref_formula(f):
                    continue
                v = cell.value
                if v is None:
                    continue
                if isinstance(v, str) and v.strip() == "":
                    continue
                row_empty = False
                fv = format_value(v)
                row_dict[h] = fv
                t = type_name(v)
                col_types.setdefault(h, []).append(t)
                if h not in col_samples:
                    col_samples[h] = fv
                if is_error(v):
                    errors.append(
                        f"Лист '{sheet_name}', ячейка "
                        f"{col_letter(ci + 1)}{cell.row}: ошибка '{v}'"
                    )
            if not row_empty:
                data_rows.append(row_dict)

        result[sheet_name] = data_rows

        if want_report:
            report_lines.append(f"### {sheet_name} — {len(data_rows)} строк")
            report_lines.append("")
            report_lines.append("| Колонка | Заголовок | Тип | Пример |")
            report_lines.append("|---|---|---|---|")
            for ci in range(n_cols):
                letter = col_letter(ci + 1)
                h = headers[ci]
                if h is None:
                    continue
                types = col_types.get(h, [])
                if not types:
                    continue
                uniq = sorted(set(types))
                tstr = uniq[0] if len(uniq) == 1 else "mixed(" + ",".join(uniq) + ")"
                sample = col_samples.get(h)
                report_lines.append(
                    f"| {letter} | {md_escape(h)} | {tstr} | "
                    f"{md_escape(short_sample(sample))} |"
                )
            report_lines.append("")

    return result, errors, "\n".join(report_lines)


def save_json(result, xlsx_path, errors_count=0):
    """
    Сохраняет '<имя_xlsx>.json' рядом с исходником.
    В начало пишется ключ _meta со служебной информацией.
    """
    json_path = xlsx_path + ".json"

    source_name = os.path.basename(xlsx_path)
    try:
        st = os.stat(xlsx_path)
        source_mtime = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        source_size = st.st_size
    except OSError:
        source_mtime = None
        source_size = None

    total_rows = sum(len(v) for v in result.values())

    meta = {
        "source_file": source_name,
        "source_mtime": source_mtime,
        "source_size_bytes": source_size,
        "converted_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "converter": CONVERTER_NAME,
        "converter_version": CONVERTER_VERSION,
        "sheets_count": len(result),
        "sheets": list(result.keys()),
        "total_rows": total_rows,
        "errors_count": errors_count,
    }

    ordered = {"_meta": meta}
    for k, v in result.items():
        key = k if k != "_meta" else "_meta_sheet"
        ordered[key] = v

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(ordered, f, ensure_ascii=False, indent=2)

    return json_path


def main():
    if len(sys.argv) < 2:
        print("Использование: python xlsx2json.py file.xlsx")
        sys.exit(1)
    path = sys.argv[1]
    if not os.path.isfile(path):
        print(f"Файл не найден: {path}")
        sys.exit(1)
    print(f"Конвертирую: {path}")
    result, errors, report = convert(path, want_report=True)
    json_path = save_json(result, path, errors_count=len(errors))
    print(f"JSON записан: {json_path}\n")
    print(report)
    if errors:
        print("=== Ошибки и замечания ===")
        for e in errors:
            print(" -", e)
    else:
        print("Ошибок нет.")


if __name__ == "__main__":
    main()